"""Create the primary Chinese Excel report for daily ETF analysis."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .market_regime import MarketRegime
from .opportunities import LANE_LABELS, STRATEGY_GROUPS, STRATEGY_LANES
from .presentation import STATUS_NAMES, STRATEGY_NAMES, chinese_frame, strategy_list
from .theme import build_theme_summary


SHEET_NAMES = (
    "今日总览", "趋势机会", "反弹机会", "防守机会", "新ETF观察", "持仓跟踪", "卖出信号",
    "策略验证", "ETF筛选过程", "风险检查", "数据异常检查", "完整候选池",
)


def _column(frame: pd.DataFrame, name: str, default=np.nan):
    return frame[name] if name in frame else pd.Series([default] * len(frame), index=frame.index)


def _lane_sheet(frame: pd.DataFrame, lane: str) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=["ETF代码", "ETF名称", "行动建议", "通俗解释"])
    common = pd.DataFrame({
        "ETF代码": frame.symbol, "ETF名称": _column(frame, "name", ""),
        "ETF类型": _column(frame, "etf_type", ""),
        "ETF结构风险": _column(frame, "structural_risk_level", ""),
        "结构风险说明": _column(frame, "structural_risk_reason", ""),
        "当前交易风险": _column(frame, "trading_risk_level", ""),
        "交易风险说明": _column(frame, "trading_risk_reason", ""),
        "行动建议": frame.action, "通俗解释": frame.action_reason,
    })
    if lane == "trend":
        extra = pd.DataFrame({
            "有效策略支持": _column(frame, "effective_strategy_consensus", ""), "有效家族支持": _column(frame, "effective_family_consensus", ""), "全部趋势家族覆盖": _column(frame, "total_family_coverage", ""),
            "共识标签": _column(frame, "consensus_label", ""),
            "支持策略": frame.validated_trend_strategies.map(strategy_list), "风险收益表现": _column(frame, "trend_mean_sharpe"),
            "近5日涨跌幅": frame.return_5d, "近20日涨跌幅": frame.return_20d, "MA20": frame.ma20, "MA60": frame.ma60,
            "距离MA20": frame.distance_ma20, "距离20日高点": frame.distance_20d_high, "波动率": frame.volatility_20d,
        })
    elif lane == "rebound":
        extra = pd.DataFrame({
            "有效策略支持": _column(frame, "effective_strategy_consensus", ""), "有效家族支持": _column(frame, "effective_family_consensus", ""), "全部反弹家族覆盖": _column(frame, "total_family_coverage", ""),
            "共识标签": _column(frame, "consensus_label", ""),
            "支持策略": frame.validated_rebound_strategies.map(strategy_list), "止跌确认分": frame.stopping_score,
            "弱优势家族支持": _column(frame, "weak_rebound_support", 0),
            "弱优势策略": _column(frame, "weak_rebound_strategies", "").map(strategy_list),
            "止跌信号": frame.stopping_signals, "持续性检查": _column(frame, "rebound_persistence_reason", "暂无可靠数据"),
            "持续性是否通过": _column(frame, "rebound_persistence_pass", False).map({True: "是", False: "否"}),
            "近5日涨跌幅": frame.return_5d, "近20日涨跌幅": frame.return_20d,
            "RSI": frame.rsi14, "MA5": frame.ma5, "MA20": frame.ma20, "重新站上MA5": frame.above_ma5.map({True: "是", False: "否"}),
            "最近3日未创新低": frame.recent_3d_no_new_low.map({True: "是", False: "否"}),
            "历史信号后10日平均收益": _column(frame, "rebound_forward_10d_mean"),
            "历史信号后10日上涨概率": _column(frame, "rebound_forward_10d_win_rate"),
        })
    else:
        extra = pd.DataFrame({
            "有效策略支持": _column(frame, "effective_strategy_consensus", ""), "有效家族支持": _column(frame, "effective_family_consensus", ""), "全部防守家族覆盖": _column(frame, "total_family_coverage", ""),
            "共识标签": _column(frame, "consensus_label", ""),
            "支持策略": frame.validated_defense_strategies.map(strategy_list), "风险收益表现": _column(frame, "defense_mean_sharpe"),
            "波动率": frame.volatility_20d, "最大回撤": frame.max_drawdown_60d, "近20日收益": frame.return_20d,
        })
    return pd.concat([common.iloc[:, :2], extra, common.iloc[:, 2:]], axis=1)


def _new_etf_sheet(frame: pd.DataFrame | None) -> pd.DataFrame:
    if frame is None or frame.empty:
        return pd.DataFrame([{"说明": "当前没有符合基础规模和流动性条件的新ETF"}])
    return pd.DataFrame({
        "ETF代码": frame["symbol"], "ETF名称": _column(frame, "name", ""),
        "上市资格": _column(frame, "qualification", ""), "上市天数": _column(frame, "listing_days"),
        "有效行情天数": _column(frame, "history_days"), "20日成交额中位数": _column(frame, "median_amount"),
        "20日平均成交额": _column(frame, "avg_amount"), "基金规模": _column(frame, "aum_yuan"),
        "近5日涨跌幅": _column(frame, "return_5d"), "近20日涨跌幅": _column(frame, "return_20d"),
        "MA20": _column(frame, "ma20"), "距离MA20": _column(frame, "distance_ma20"),
        "ETF结构风险": _column(frame, "structural_risk_level", "暂无可靠数据"),
        "结构风险说明": _column(frame, "structural_risk_reason", "暂无可靠数据"),
        "状态": _column(frame, "observation_status", "仅观察"),
        "说明": _column(frame, "observation_explanation", "新ETF不参与正式历史验证，也不产生买入建议。"),
    })


def _holding_sheet(frame: pd.DataFrame | None) -> pd.DataFrame:
    if frame is None or frame.empty:
        return pd.DataFrame([{"说明": "当前未登记真实持仓；可通过 --holdings 指定人工持仓文件。未提供买入日期和成本时，系统不会生成个性化卖出建议。"}])
    return pd.DataFrame({
        "ETF代码": frame["symbol"], "ETF名称": _column(frame, "name", ""),
        "持仓来源": _column(frame, "holding_source", ""), "机会类型": _column(frame, "opportunity_type", ""),
        "买入日期": _column(frame, "buy_date", ""), "买入价": _column(frame, "buy_price"),
        "当前价": _column(frame, "current_price"), "持仓数量": _column(frame, "shares"),
        "持仓成本": _column(frame, "buy_price") * _column(frame, "shares"),
        "当前收益": _column(frame, "current_return"), "持仓最高收益": _column(frame, "peak_return"),
        "从最高点回撤": _column(frame, "drawdown_from_peak"), "MA5": _column(frame, "ma5"),
        "MA20": _column(frame, "ma20"), "MA60": _column(frame, "ma60"),
        "当前策略支持": _column(frame, "current_strategy_support", 0),
        "买入时策略支持": _column(frame, "entry_strategy_support", 0),
        "止跌确认分（反弹）": _column(frame, "stopping_score"),
        "当前交易风险": _column(frame, "current_trading_risk", "暂无可靠数据"),
        "交易风险说明": _column(frame, "current_trading_risk_reason", "暂无可靠数据"),
        "买入时市场环境": _column(frame, "entry_market_regime", "未记录"),
        "买入理由": _column(frame, "entry_reason", "未记录"),
        "买入理由是否仍成立": _column(frame, "entry_reason_still_valid", False).map({True: "是", False: "否"}),
        "持仓状态": _column(frame, "holding_status", ""), "卖出建议": _column(frame, "exit_advice", ""),
        "原因": _column(frame, "exit_reason", ""),
    })


def opportunity_flags(lanes: dict[str, pd.DataFrame]) -> tuple[bool, bool]:
    """Return (worth following, executable buy) under explicit action semantics."""
    executable_actions = {"可以关注买入", "小仓反弹试错", "较高优先级反弹候选", "高优先级反弹候选"}
    worthwhile = any((frame.action != "暂不参与").any() for frame in lanes.values() if not frame.empty)
    executable = any(frame.action.isin(executable_actions).any() for frame in lanes.values() if not frame.empty)
    return worthwhile, executable


def _sell_signal_sheet(holding_tracking: pd.DataFrame | None, sell_signals: pd.DataFrame | None) -> pd.DataFrame:
    if holding_tracking is None or holding_tracking.empty:
        return pd.DataFrame([{"说明": "由于当前没有真实持仓数据，系统不生成个性化卖出建议。"}])
    if sell_signals is None or sell_signals.empty:
        return pd.DataFrame([{"说明": f"已检查 {len(holding_tracking)} 只已登记持仓，今日未触发退出条件。"}])
    return _holding_sheet(sell_signals)


def _strategy_validation(status: pd.DataFrame, backtests: pd.DataFrame, rebound: pd.DataFrame) -> pd.DataFrame:
    result = status.copy()
    result["lane"] = result.strategy.map(STRATEGY_LANES)
    result["策略类型"] = result.lane.map(LANE_LABELS).fillna("控制器或不参与比较")
    long_fields = [column for column in (
        "strategy", "passed", "sharpe", "max_drawdown", "annualized_return", "win_rate", "trades", "days",
        "one_way_cost", "slippage", "benchmark_return", "excess_return", "backtest_start_date", "backtest_end_date",
        "price_adjustment", "validation_confidence", "reason",
    ) if column in backtests]
    result = result.merge(backtests[long_fields], on="strategy", how="left", suffixes=("", "_long"))
    rebound_fields = [column for column in ("strategy", "passed", "validation_level", "validation_explanation", "signal_count", "forward_5d_mean", "forward_10d_mean", "forward_10d_win_rate", "average_adverse_excursion", "payoff_ratio", "days", "validation_confidence", "backtest_start_date", "backtest_end_date", "price_adjustment", "reason") if column in rebound]
    result = result.merge(rebound[rebound_fields], on="strategy", how="left", suffixes=("_long", "_rebound"))
    is_rebound = result.lane == "rebound"
    result["最终是否通过"] = np.where(is_rebound, result.get("passed_rebound"), result.get("passed_long"))
    output = pd.DataFrame({
        "策略中文名": result.strategy.map(lambda value: STRATEGY_NAMES.get(value, value)),
        "策略家族": result.strategy.map(STRATEGY_GROUPS).fillna(result.family.map({"not_comparable": "不参与横向比较", "market_controller": "市场环境控制", "regime": "市场状态"}).fillna(result.family)),
        "策略类型": result["策略类型"], "运行状态": result.status.map(lambda value: STATUS_NAMES.get(value, value)),
        "历史验证等级": np.where(is_rebound, result.get("validation_level", "未通过"), result["最终是否通过"].map({True: "合格", False: "未通过"}).fillna("不适用")),
        "风险收益表现": result.get("sharpe"), "最大回撤": result.get("max_drawdown"), "年化收益": result.get("annualized_return"),
        "胜率（日收益口径）": result.get("win_rate"),
        "交易或信号数量": np.where(is_rebound, result.get("signal_count"), result.get("trades")),
        "有效交易日数量": np.where(is_rebound, result.get("days_rebound"), result.get("days_long")),
        "历史验证置信度": np.where(is_rebound, result.get("validation_confidence_rebound"), result.get("validation_confidence_long")),
        "5日后平均收益": result.get("forward_5d_mean"), "10日后平均收益": result.get("forward_10d_mean"),
        "10日上涨概率": result.get("forward_10d_win_rate"), "平均最大不利波动": result.get("average_adverse_excursion"),
        "盈亏比": result.get("payoff_ratio"),
        "回测起始日期": np.where(is_rebound, result.get("backtest_start_date_rebound"), result.get("backtest_start_date_long")),
        "回测结束日期": np.where(is_rebound, result.get("backtest_end_date_rebound"), result.get("backtest_end_date_long")),
        "手续费假设（单边）": result.get("one_way_cost"),
        "滑点假设": result.get("slippage"),
        "基准收益": result.get("benchmark_return"),
        "相对基准表现": result.get("excess_return"),
        "复权数据说明": np.where(is_rebound, result.get("price_adjustment_rebound"), result.get("price_adjustment_long")),
        "备注": np.where(
            is_rebound,
            "反弹策略按短期信号结果分级；弱优势不计入有效策略共识。" + result.get("validation_explanation", "").fillna("").map(lambda value: f" {value}" if value else "") + result.get("reason_rebound", "").fillna("").map(lambda value: f" {value}" if value else ""),
            result.get("reason_long", ""),
        ),
    })
    return output


def _format_workbook(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    action_fills = {
        "可以关注买入": "C6EFCE", "小仓反弹试错": "C6EFCE", "较高优先级反弹候选": "C6EFCE",
        "高优先级反弹候选": "C6EFCE", "稳健持有候选": "C6EFCE",
        "等待回调": "FFF2CC", "继续观察": "FFF2CC",
        "暂不追高": "FCE4D6", "暂不参与": "E7E6E6", "高": "F4CCCC",
    }
    percent_words = ("涨跌幅", "年化收益", "日后平均收益", "日平均收益", "近20日收益", "波动率", "回撤", "距离", "上涨概率", "胜率", "占比", "折溢价", "基础动量")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_index, cells in enumerate(sheet.columns, 1):
            values = [str(cell.value) if cell.value is not None else "" for cell in cells[:100]]
            width = min(45, max(10, max((len(value) for value in values), default=8) + 2))
            sheet.column_dimensions[get_column_letter(column_index)].width = width
            header = str(cells[0].value or "")
            for cell in cells[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=len(str(cell.value or "")) > 25)
                if header in ("风险收益表现", "盈亏比") and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00"
                elif any(word in header for word in percent_words) and isinstance(cell.value, (int, float)):
                    cell.number_format = "0.00%"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.00"
        for row in range(2, sheet.max_row + 1):
            for cell in sheet[row]:
                if str(cell.value) in action_fills:
                    cell.fill = PatternFill("solid", fgColor=action_fills[str(cell.value)])
                    cell.font = Font(bold=True)
        if sheet.title == "今日总览":
            percentage_items = {
                "近5日收益离散度", "近20日收益离散度",
                "强弱两端20日收益差", "近20日上涨ETF占比",
            }
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value in percentage_items:
                    sheet.cell(row, 2).number_format = "0.00%"
    workbook.save(path)


def generate_excel_report(
    output_dir: Path,
    report_date: date,
    as_of: str,
    regime: MarketRegime,
    summary: dict,
    lanes: dict[str, pd.DataFrame],
    strategy_status: pd.DataFrame,
    backtests: pd.DataFrame,
    rebound_validation: pd.DataFrame,
    audit: pd.DataFrame,
    candidates: pd.DataFrame,
    risks: pd.DataFrame,
    shortlist: pd.DataFrame,
    summary_lanes: dict[str, pd.DataFrame] | None = None,
    price_anomalies: pd.DataFrame | None = None,
    new_etfs: pd.DataFrame | None = None,
    holding_tracking: pd.DataFrame | None = None,
    sell_signals: pd.DataFrame | None = None,
) -> Path:
    """Write the Chinese multi-sheet report and return the workbook path."""
    active = {lane: int((frame.action != "暂不参与").sum()) if not frame.empty else 0 for lane, frame in lanes.items()}
    worthwhile, executable = opportunity_flags(lanes)
    duplicate_mask = audit.reasons.fillna("").str.contains("跟踪同一指数|相似主题且收益相关性过高")
    basic_removed = int((~audit.eligible & ~duplicate_mask).sum())
    after_basic = len(audit) - basic_removed
    display_lanes = summary_lanes or lanes
    theme_summary = build_theme_summary(display_lanes)
    trend_names = "、".join(display_lanes["trend"].loc[display_lanes["trend"].action != "暂不参与", "name"].head(2))
    rebound_names = "、".join(display_lanes["rebound"].loc[display_lanes["rebound"].action != "暂不参与", "name"].head(2))
    sentence = f"{regime.label}。"
    sentence += f"趋势方向有{active['trend']}只候选" + (f"，包括{trend_names}" if trend_names else "") + "；"
    sentence += f"反弹方向有{active['rebound']}只候选" + (f"，包括{rebound_names}" if rebound_names else "") + f"；防守方向有{active['defense']}只候选。"
    if not executable:
        sentence += "当前有可继续跟踪的方向，但没有达到系统定义的可执行买入门槛。" if worthwhile else "当前没有达到关注门槛的机会。"
    theme_text = {}
    for lane, label in (("trend", "趋势"), ("rebound", "反弹"), ("defense", "防守")):
        rows = theme_summary[theme_summary.opportunity_lane == lane] if not theme_summary.empty else pd.DataFrame()
        theme_text[lane] = "；".join(
            f"{row.theme}：代表 {row.representative_symbol} {row.representative_name}，备选 {row.alternative_etfs}"
            for _, row in rows.iterrows()
        ) or "暂无"
    overview = pd.DataFrame({"项目": [
        "报告生成日期", "行情截止日期", "大盘状态", "市场结构", "当前市场", "大盘状态说明", "市场结构说明", "全市场ETF数量", "基础筛选后数量", "同类去重后数量",
        "有效策略数量", "当前趋势候选数量", "当前反弹候选数量", "当前防守候选数量", "新ETF观察数量", "当前持仓数量", "今日卖出信号数量",
        "候选ETF横截面数量", "近5日收益离散度", "近20日收益离散度", "强弱两端20日收益差", "近20日上涨ETF占比",
        "今日存在值得关注的机会", "今日存在可执行买入信号", "趋势主题摘要", "反弹主题摘要", "防守主题摘要", "今日一句话总结",
    ], "内容": [
        report_date.isoformat(), as_of, regime.broad_label, regime.structure_label, regime.label, regime.broad_explanation, regime.structure_explanation, summary.get("universe_count", 0), after_basic, summary.get("candidate_count", 0),
        summary.get("validated_strategy_count", 0), active["trend"], active["rebound"], active["defense"], summary.get("new_etf_count", 0), summary.get("holding_count", 0), summary.get("sell_signal_count", 0),
        regime.cross_section_count, regime.return_5d_std, regime.return_20d_std,
        regime.top_bottom_20d_spread, regime.advancing_ratio_20d, "是" if worthwhile else "否", "是" if executable else "否",
        theme_text["trend"], theme_text["rebound"], theme_text["defense"],
        sentence,
    ]})
    process = pd.DataFrame([
        {"阶段": "初始全市场", "数量": len(audit), "说明": "本地数据库可识别ETF"},
        {"阶段": "新ETF观察池", "数量": int(audit.get("observation_eligible", False).sum()), "说明": "上市不足180天，只观察，不参加正式历史验证或产生买入建议"},
        {"阶段": "基础条件淘汰", "数量": int((~audit.eligible & ~duplicate_mask & ~audit.get("observation_eligible", False)).sum()), "说明": "类型、规模、成交额或资料不合格"},
        {"阶段": "同类去重淘汰", "数量": int(duplicate_mask.sum()), "说明": "同指数或收益高度相关时保留更活跃品种"},
        {"阶段": "最终可比较候选", "数量": len(candidates), "说明": "进入三赛道策略比较"},
    ])
    eliminated = chinese_frame(audit[~audit.eligible].copy())
    eliminated.insert(0, "阶段", "淘汰明细")
    process = pd.concat([process, eliminated], ignore_index=True, sort=False)
    pool = candidates.copy()
    flags = shortlist.set_index("symbol")[[column for column in shortlist if column.startswith("in_")]] if not shortlist.empty else pd.DataFrame()
    pool = pool.merge(flags, left_on="symbol", right_index=True, how="left")
    for lane, label in LANE_LABELS.items():
        pool[f"是否进入{label}"] = pool.get(f"in_{lane}_shortlist", False).fillna(False).map({True: "是", False: "否"})
    pool = chinese_frame(pool.drop(columns=[column for column in pool if column.startswith("in_")], errors="ignore"))

    valid_defense = int(summary.get("validated_lane_denominators", {}).get("defense", 0))
    overview = pd.concat([overview, pd.DataFrame({"项目": ["防守赛道状态"], "内容": [
        "当前暂无有效防守策略，系统暂时无法可靠评价防守机会"
        if valid_defense == 0 else ("今日暂无满足条件的防守机会" if active["defense"] == 0 else "防守模型有效，今日存在候选")
    ]})], ignore_index=True)
    risk_columns = [column for column in (
        "symbol", "name", "etf_type", "liquidity_status", "liquidity_reason", "premium_status", "premium_rate",
        "premium_reason", "special_risk_note", "price_risk_status", "price_risk_reason", "announcement_status",
        "announcement_warnings", "structural_risk_level", "structural_risk_reason", "warnings",
    ) if column in risks]
    risk_sheet = risks[risk_columns].copy()
    action_risks = pd.concat([frame for frame in lanes.values() if not frame.empty], ignore_index=True, sort=False)
    if not action_risks.empty:
        trade = action_risks[["symbol", "trading_risk_level", "trading_risk_reason"]].drop_duplicates("symbol")
        risk_sheet = risk_sheet.merge(trade, on="symbol", how="left")
    anomaly_sheet = (price_anomalies.copy() if price_anomalies is not None else pd.DataFrame())
    if anomaly_sheet.empty:
        anomaly_sheet = pd.DataFrame([{"备注": "未发现明显价格断层"}])
    else:
        names = pd.concat([
            audit[["symbol", "name"]], candidates[["symbol", "name"]], risks[["symbol", "name"]]
        ], ignore_index=True).drop_duplicates("symbol").set_index("symbol")["name"].to_dict()
        anomaly_sheet["name"] = anomaly_sheet["symbol"].map(names).fillna("")
        anomaly_sheet = anomaly_sheet.rename(columns={
            "symbol": "ETF代码", "name": "ETF名称", "date": "日期", "previous_close": "前一日价格",
            "current_close": "当前价格", "daily_change": "单日变化幅度", "suspected_corporate_action": "是否疑似拆分/除权",
            "source": "数据源", "repaired": "是否已修复", "repair_method": "修复方式", "note": "备注",
            "repair_confidence": "修正可信度", "confirmation_source": "确认来源",
            "manual_review_required": "是否需要人工核验",
        })
        for column in ("是否疑似拆分/除权", "是否已修复", "是否需要人工核验"):
            anomaly_sheet[column] = anomaly_sheet[column].map(lambda value: "是" if bool(value) else "否")
        anomaly_sheet = anomaly_sheet[["ETF代码", "ETF名称", "日期", "前一日价格", "当前价格", "单日变化幅度", "是否疑似拆分/除权", "数据源", "是否已修复", "修复方式", "修正可信度", "确认来源", "是否需要人工核验", "备注"]]
    observation_sheet = _new_etf_sheet(new_etfs)
    holding_sheet = _holding_sheet(holding_tracking)
    sell_sheet = _sell_signal_sheet(holding_tracking, sell_signals)
    sheets = {
        "今日总览": overview, "趋势机会": _lane_sheet(lanes["trend"], "trend"),
        "反弹机会": _lane_sheet(lanes["rebound"], "rebound"), "防守机会": _lane_sheet(lanes["defense"], "defense"),
        "新ETF观察": observation_sheet, "持仓跟踪": holding_sheet, "卖出信号": sell_sheet,
        "策略验证": _strategy_validation(strategy_status, backtests, rebound_validation), "ETF筛选过程": process,
        "风险检查": chinese_frame(risk_sheet), "数据异常检查": anomaly_sheet, "完整候选池": pool,
    }
    path = output_dir / f"ETF每日分析_{report_date:%Y%m%d}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name in SHEET_NAMES:
            sheets[name].to_excel(writer, sheet_name=name, index=False)
    _format_workbook(path)
    return path
