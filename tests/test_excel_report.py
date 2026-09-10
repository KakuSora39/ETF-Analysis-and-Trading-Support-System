import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from etf_universe.excel_report import SHEET_NAMES, _format_workbook, _sell_signal_sheet, _strategy_validation, opportunity_flags
from etf_universe.presentation import STRATEGY_NAMES


class ExcelReportTests(unittest.TestCase):
    def test_waiting_only_is_worth_following_but_not_executable(self):
        lanes = {
            "trend": pd.DataFrame([{"action": "等待回调"}]),
            "rebound": pd.DataFrame(),
            "defense": pd.DataFrame([{"action": "防守候选"}]),
        }
        self.assertEqual(opportunity_flags(lanes), (True, False))

    def test_no_positions_does_not_claim_no_sell_signal(self):
        result = _sell_signal_sheet(pd.DataFrame(), pd.DataFrame())
        self.assertIn("没有真实持仓数据", result.iloc[0, 0])
        self.assertIn("不生成个性化卖出建议", result.iloc[0, 0])
    def test_report_has_new_etf_and_holding_sheets(self):
        self.assertEqual(len(SHEET_NAMES), 12)
        self.assertIn("新ETF观察", SHEET_NAMES)
        self.assertIn("持仓跟踪", SHEET_NAMES)
        self.assertIn("卖出信号", SHEET_NAMES)

    def test_strategy_validation_keeps_days_and_confidence_by_lane(self):
        status = pd.DataFrame({
            "strategy": ["momentum_rotation", "mean_reversion"], "family": ["趋势", "反弹"],
            "status": ["voted", "voted"],
        })
        backtests = pd.DataFrame({
            "strategy": ["momentum_rotation"], "passed": [True], "days": [93],
            "validation_confidence": ["低"], "reason": [""],
        })
        rebound = pd.DataFrame({
            "strategy": ["mean_reversion"], "passed": [True], "validation_level": ["合格"],
            "validation_explanation": ["达到短期门槛"], "days": [180],
            "validation_confidence": ["中"], "reason": [""],
        })
        result = _strategy_validation(status, backtests, rebound)
        values = result.set_index("策略中文名")
        trend_name = STRATEGY_NAMES.get("momentum_rotation", "momentum_rotation")
        rebound_name = STRATEGY_NAMES.get("mean_reversion", "mean_reversion")
        self.assertEqual(values.loc[trend_name, "有效交易日数量"], 93)
        self.assertEqual(values.loc[trend_name, "历史验证置信度"], "低")
        self.assertEqual(values.loc[rebound_name, "有效交易日数量"], 180)
        self.assertEqual(values.loc[rebound_name, "历史验证置信度"], "中")
    def test_readable_headers_percentages_and_actions(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "测试报告.xlsx"
            pd.DataFrame({
                "风险收益表现": [1.08], "历史信号后10日平均收益": [.04],
                "行动建议": ["小仓反弹试错"],
            }).to_excel(path, index=False, sheet_name="反弹机会")
            _format_workbook(path)
            sheet = load_workbook(path)["反弹机会"]
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:C2")
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(sheet["A2"].number_format, "0.00")
            self.assertEqual(sheet["B2"].number_format, "0.00%")
            self.assertTrue(sheet["C2"].font.bold)

    def test_overview_dispersion_metrics_are_percentages(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "测试报告.xlsx"
            pd.DataFrame({
                "项目": ["近5日收益离散度", "近20日上涨ETF占比"],
                "内容": [.038, .42],
            }).to_excel(path, index=False, sheet_name="今日总览")
            _format_workbook(path)
            sheet = load_workbook(path)["今日总览"]
            self.assertEqual(sheet["B2"].number_format, "0.00%")
            self.assertEqual(sheet["B3"].number_format, "0.00%")


if __name__ == "__main__":
    unittest.main()
