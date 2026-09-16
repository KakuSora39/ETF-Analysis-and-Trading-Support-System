import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from simulation.etf_universe.reporter import append_excel


class PaperExcelTests(unittest.TestCase):
    def test_sheets_and_no_trade_message(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "report.xlsx"
            Workbook().save(path)
            states = {}
            for mode in ("conservative", "progressive"):
                states[mode] = {"mode": mode, "strategy_version": mode + "-v1",
                                "initial_capital": 100000, "cash": 100000, "positions": {},
                                "nav_history": [], "trade_log": []}
            append_excel(path, states, {}, "2026-09-16", {})
            book = load_workbook(path, read_only=True)
            self.assertIn("模拟盘总览", book.sheetnames)
            self.assertIn("A_B模拟对比", book.sheetnames)
            self.assertEqual("今日无模拟交易", book["模拟交易记录"]["A2"].value)
            book.close()


if __name__ == "__main__":
    unittest.main()
